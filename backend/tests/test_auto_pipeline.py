"""上传即评审：解析 → 规则提取 → 冻结输入 → 核对与预测自动衔接，并产出报告、响应矩阵、报价与度量（真实 PG、受限 worker 角色、脚本化模型）。"""

from __future__ import annotations

import io
import json
from typing import Any

import pypdfium2 as pdfium
import pytest

from bidwb import tasks
from bidwb.ai import features
from bidwb.ai.handlers import build_handlers
from bidwb.config import Settings
from bidwb.messaging import InMemoryBus
from bidwb.workers.runner import WorkerRunner
from tests.fakes import ScriptedLlm, docs
from tests.pdfs import make_pdf
from tests.test_files_and_parse import _setup, _upload, file_id_of

pytestmark = pytest.mark.integration

TENDER = make_pdf(
    [
        ["第一章 工期要求", "总工期不得超过一百八十日历天，超过的按废标处理。"],
        ["第二章 质量要求", "工程质量必须达到合格标准。"],
        ["第三章 评标办法", "施工进度计划合理可行的得8分，计划不合理的得0分。"],
        ["第四章 评分构成", "技术标满分为60分。", "价格分满分30分，高于基准价每1%扣1分，低于基准价每1%扣0.5分。"],
    ]
)
BODY = make_pdf(
    [
        ["第一章 施工进度计划", "本工程计划总工期为二百四十日历天。"],
        ["第二章 质量保证", "本工程质量目标为合格，满足设计及验收规范。"],
        ["第三章 施工方法", "采用沉井施工，沉井下沉到位后封底。最大开挖深度为12m。"],
        ["第四章 基坑支护", "基坑支护方案按最大开挖深度为15m 设计。"],
    ]
)
BODY_V2 = make_pdf(
    [
        ["第一章 施工进度计划", "本工程计划总工期为一百七十日历天。"],
        ["第二章 质量保证", "本工程质量目标为合格，满足设计及验收规范。"],
    ]
)
DEADLINE = "总工期不得超过一百八十日历天"
QUALITY = "工程质量必须达到合格标准"
SCORING = "施工进度计划合理可行的得8分"
FULL_SCORE = "技术标满分为60分"
PRICE_RULE = "价格分满分30分，高于基准价每1%扣1分，低于基准价每1%扣0.5分"
LATE = "计划总工期为二百四十日历天"
GOOD = "本工程质量目标为合格"


@pytest.fixture
def settings() -> Settings:
    return Settings(
        environment="test", cookie_secure=False, session_cookie_name="sid", login_rate_limit_per_minute=1000, auto_analysis=True
    )


def mapping_reply(user: str) -> dict[str, Any]:
    """条款对应章节：评分条款对到标题含“进度”的章节。"""
    items = docs(user)
    sections = [i for i, _ in items if i.startswith("S")]
    progress = [i for i, text in items if i.startswith("S") and "进度" in text.split("｜")[0]]
    rules = [i for i, _ in items if i.startswith("R")]
    return {"mapping": [{"item": r, "sections": progress[:1] if "进度" in user else sections[:1]} for r in rules]}


def judgement_reply(user: str) -> dict[str, Any]:
    """逐项判断：三档里判为“良”（2～6 分）、给 4 分；写到了总工期，没写关键线路说明。"""
    body = next((i for i in docs(user) if "总工期" in i[1] and "本工程" in i[1]), None)
    return {
        "grades": [
            {"label": "优", "low": "6.5", "high": "8"},
            {"label": "良", "low": "2", "high": "6"},
            {"label": "差", "low": "0", "high": "1.5"},
        ],
        "grade": 1,
        "score": "4",
        "key_points": [
            {"point": "总工期安排", "status": "covered", "evidence_ids": [body[0]] if body else []},
            {"point": "关键线路说明", "status": "missing", "evidence_ids": []},
        ],
        "specificity": 0.6,
        "effectiveness": {"status": "weak", "note": "缺少关键线路", "evidence_ids": [body[0]] if body else []},
        "consistency": {"status": "unclear", "note": "未检查到双侧证据", "evidence_ids": []},
        "specificity_note": "写明了本工程的计划工期",
        "assessment": "进度计划给出了总工期但缺少关键线路说明",
        "limitations": "评分标准含主观判断",
    }


def _respond(system: str, user: str) -> Any:  # noqa: PLR0911  按提示词分派的测试模型，每类提示词一个出口
    items = docs(user)
    if "抽取员" in system:
        out = []
        for kind, statement, quote, points in (
            ("disqualification", "总工期不超过180日历天", DEADLINE, None),
            ("mandatory", "工程质量必须达到合格标准", QUALITY, None),
            ("scoring", "施工进度计划合理可行", SCORING, "8"),
        ):
            hit = next((i for i in items if quote in i[1]), None)
            if hit:
                out.append({"kind": kind, "statement": statement, "evidence_ids": [hit[0]], "quote": quote, "max_points": points})
        terms: dict[str, Any] = {}
        full = next((i for i in items if FULL_SCORE in i[1]), None)
        if full:
            terms["technical_full_score"] = {"value": "60", "evidence_ids": [full[0]], "quote": FULL_SCORE}
        price = next((i for i in items if PRICE_RULE in i[1]), None)
        if price:  # 最高限价 99999 在原文里找不到：只丢这一个数值，其余保留
            terms["price_rule"] = {
                "full_score": "30",
                "above_deduct_per_pct": "1",
                "below_deduct_per_pct": "0.5",
                "max_limit": "99999",
                "evidence_ids": [price[0]],
                "quote": PRICE_RULE,
            }
        return {"rules": out, **terms}
    if "目录分析员" in system:
        return mapping_reply(user)
    if "评标专家" in system:
        return judgement_reply(user)
    statement = user.split("）：", maxsplit=1)[1].split("\n", maxsplit=1)[0]
    if statement.startswith("总工期"):
        late = next((i for i in items if LATE in i[1]), None)
        if late is None:
            good = next((i for i in items if "计划总工期" in i[1]), None)
            return {
                "verdict": "satisfied",
                "verdict_evidence_ids": [good[0]] if good else [],
                "verdict_note": "计划工期未超过上限",
                "findings": [],
            }
        return {
            "verdict": "violated",
            "verdict_note": "计划工期超过上限",
            "findings": [
                {
                    "nature": "fact_conflict",
                    "severity": "critical",
                    "statement": "进度计划工期超过招标上限，可能被否决",
                    "body_evidence_ids": [late[0]],
                    "rule_evidence_ids": [items[0][0]],
                    "body_quote": LATE,
                    "rule_quote": DEADLINE,
                    "required_facts": [],
                }
            ],
        }
    if statement.startswith("工程质量"):
        good = next((i for i in items if GOOD in i[1]), None)
        return {
            "verdict": "satisfied",
            "verdict_evidence_ids": [good[0]] if good else [],
            "verdict_note": "质量目标满足要求",
            "findings": [],
        }
    return {"verdict": "unclear", "verdict_note": "评分条款由预测评估", "findings": []}


def _drain(parse_runner: WorkerRunner, ai_runner: WorkerRunner, bus: InMemoryBus, engine) -> None:  # type: ignore[no-untyped-def]
    """反复转发发件箱并消费，直到流水线停下（每个任务完成后会自动推进并入队下一步）。"""
    for _ in range(30):
        with engine.begin() as conn:
            relayed = tasks.relay_outbox(conn, bus.producer())
        worked = False
        while parse_runner.poll_once() or ai_runner.poll_once():
            worked = True
        if not relayed and not worked:
            return
    raise AssertionError("流水线没有收敛")


class Env:
    def __init__(self, api, pid, store, engine, bus, parse_runner, ai_runner, make_member, api_factory) -> None:  # type: ignore[no-untyped-def]
        self.api, self.pid, self.store, self.engine, self.bus = api, pid, store, engine, bus
        self.parse_runner, self.ai_runner = parse_runner, ai_runner
        self.make_member, self.api_factory = make_member, api_factory

    def drain(self) -> None:
        _drain(self.parse_runner, self.ai_runner, self.bus, self.engine)

    def get(self, path: str) -> Any:
        r = self.api.get(f"/projects/{self.pid}{path}")
        assert r.status_code == 200, r.text
        return r.json()

    def state(self) -> str:
        return str(self.api.get(f"/projects/{self.pid}").json()["state"])


def _env(api_factory, make_member, store, engine, worker_container, llm) -> Env:  # type: ignore[no-untyped-def]
    api, pid = _setup(api_factory, make_member)
    bus = InMemoryBus()
    handlers = build_handlers(llm)
    parse_runner = WorkerRunner(worker_container, "parse", bus.consumer(["job.parse"]), handlers, consumer_name="worker-parse")
    ai_runner = WorkerRunner(worker_container, "ai", bus.consumer(["job.ai"]), handlers, consumer_name="worker-ai")
    return Env(api, pid, store, engine, bus, parse_runner, ai_runner, make_member, api_factory)


def test_upload_runs_the_whole_review_and_builds_the_report(api_factory, make_member, store, engine, worker_container) -> None:  # type: ignore[no-untyped-def]
    env = _env(api_factory, make_member, store, engine, worker_container, ScriptedLlm(_respond))
    _upload(env.api, store, env.pid, TENDER)
    assert env.state() == "materials_pending" and env.get("/pipeline")["missing"] == ["body"]  # 只有招标文件：等待正文
    body = _upload(env.api, store, env.pid, BODY, name="施工组织设计.pdf", role="body", key="body-upload-0001")[1].json()["version_id"]
    assert env.state() == "analyzing"
    env.drain()

    pipeline = env.get("/pipeline")
    assert pipeline["stage"] == "done" and [s["state"] for s in pipeline["steps"]] == ["done"] * 4
    assert pipeline["rules_total"] == 3 and pipeline["rules_unconfirmed"] == 3 and pipeline["round_no"] == 1
    assert env.state() == "issues_open"

    report = env.get("/report")
    matrix = report["qualification"]
    assert matrix["qualification"] == {"total": 2, "satisfied": 1, "violated": 1, "unclear": 0, "failed": 0}
    assert matrix["scoring"]["total"] == 1 and matrix["coverage"]["status"] == "complete"
    satisfied = next(r for r in matrix["rows"] if r["verdict"] == "satisfied")
    assert satisfied["kind"] == "mandatory" and satisfied["rule_basis"] == "unconfirmed" and satisfied["evidence"][0]["version_id"] == body
    by = {c["category"]: c for c in report["categories"]}
    assert (
        by["disqualification"]["open"] == 1
        and by["consistency"]["total"] == 1
        and by["scope"]["total"] == 1
        and by["response"]["total"] == 0
    )
    assert report["basis"]["rules_unconfirmed"] == 3
    prediction = report["prediction"]
    assert (
        prediction["status"] == "candidate" and prediction["purpose"] == "research" and prediction["validation_status"] == "not_validated"
    )
    assert (prediction["estimated_value"], prediction["low"], prediction["high"], prediction["maximum"]) == ("4", "2", "6", "8")
    assert prediction["calibration"]["state"] == "uncalibrated" and prediction["topics"][0]["evidence"]
    assert prediction["scoring_model"] is None and "尚未用官方成绩校准" in prediction["reason"]
    assert prediction["topics"][0]["topic"] == "施工进度计划合理可行（满分 8 分，判为“良”档；缺：关键线路说明）"
    with engine.connect() as conn:
        key = conn.exec_driver_sql(
            "SELECT a.object_key FROM predictions p JOIN artifacts a ON a.id=p.detail_artifact_id WHERE p.project_id=%s", (env.pid,)
        ).scalar()
    detail = json.loads(store.get(key))
    (item,) = detail["items"]
    assert detail["mapping"]["status"] == "model" and item["mapped"] and item["sections"] == ["第一章 施工进度计划"]
    assert item["scores"] == ["4", "4", "4"] and item["grade_label"] == "良" and item["grade_position"] == 0.5 and item["coverage"] == 0.5
    assert item["item_id"] == item["rule_id"] and "施工进度计划合理可行的得8分" in item["rubric"]
    assert detail["feature_spec"] == "pf-1" and set(detail["features"]) == set(features.FEATURES)
    assert detail["features"]["raw_pct"] == 50 and detail["features"]["max_ratio"] == round(8 / 60, 6)
    assert detail["calibrated"] is None and detail["shadow"] is None and detail["anchors"] is None
    assert [a["key"] for a in report["next_actions"]] == ["handle_disqualification", "confirm_rules", "handle_findings", "set_price"]

    findings = env.get("/findings?limit=100")["items"]
    assert {f["category"]: f["source"] for f in findings} == {"disqualification": "rule", "consistency": "consistency", "scope": "scope"}
    disq = next(f for f in findings if f["category"] == "disqualification")
    assert disq["rule_basis"] == "unconfirmed" and disq["impact"] == "qualification" and disq["rule_id"]
    assert next(f for f in findings if f["category"] == "scope")["rule_id"] is None
    assert env.get("/findings?category=scope")["items"][0]["category"] == "scope"

    # 提交闸门：废标／强制类规则未确认、关键问题未处理，都不能登记提交
    submit = {
        "snapshot_id": report["snapshot_id"],
        "body_version_id": body,
        "submitted_at": "2026-10-01T08:00:00+08:00",
        "receipt_artifact_id": None,
    }
    blocked = env.api.post(f"/projects/{env.pid}/submissions", submit)
    assert blocked.status_code == 409 and blocked.json()["code"] == "qualification_rules_unconfirmed"
    rules = env.get("/rules")
    for rule in rules["rules"]:
        if rule["kind"] in ("disqualification", "mandatory"):
            ok = env.api.post(
                f"/rules/{rule['rule_id']}/reviews", {"object_version": str(rules["revision"]), "decision": "confirm", "reason": "核对原文"}
            )
            assert ok.status_code == 200, ok.text
    assert env.get("/rules")["rule_set_id"] == rules["rule_set_id"]  # 只是确认，不改变参与评审的规则集合，不复制规则包
    assert env.get("/report")["basis"]["rules_unconfirmed"] == 1
    assert env.api.post(f"/projects/{env.pid}/submissions", submit).json()["code"] == "critical_findings_open"

    # 报价
    plan = {
        "params": {
            "full_score": "30",
            "above_deduct_per_pct": "1",
            "below_deduct_per_pct": "0.5",
            "max_limit": "10500",
            "abnormal_low_pct": "10",
            "benchmark": {"method": "average", "coefficient": "1", "trim_pct": "0", "include_own": False},
        },
        "competitor_prices": ["9800", "10000", "10200"],
        "candidate_prices": ["9900", "10300", "10600"],
        "note": "",
    }
    assert env.get("/price") is None
    saved = env.api.put(f"/projects/{env.pid}/price", plan)
    assert saved.status_code == 200, saved.text
    sim = saved.json()["simulation"]
    assert sim["benchmark"] == "10000" and sim["best"]["price"] == "9900" and sim["candidates"][2]["flags"] == ["over_max_limit"]
    preview = env.api.post(f"/projects/{env.pid}/price:simulate", {**plan, "candidate_prices": ["10000"]})
    assert preview.status_code == 200 and preview.json()["candidates"][0]["score"] == "30"
    assert len(env.get("/price/history")) == 1
    assert env.get("/report")["price"]["best"]["price"] == "9900"

    # 导出：Word 报告与 PDF 批注版
    docx = env.api.get(f"/projects/{env.pid}/report.docx")
    assert docx.status_code == 200 and docx.content[:2] == b"PK" and "attachment" in docx.headers["content-disposition"]
    annotated = env.api.get(f"/projects/{env.pid}/body.annotated.pdf")
    assert annotated.status_code == 200 and annotated.content[:5] == b"%PDF-"
    document = pdfium.PdfDocument(io.BytesIO(annotated.content))
    assert sum(pdfium.raw.FPDFPage_GetAnnotCount(page) for page in document) >= 2  # 废标风险、沉井等定位到正文页

    # 度量
    metrics = env.api.get("/metrics").json()
    assert (
        metrics["qualification"]["identified"] == 1
        and metrics["outcomes"]["submissions"] == 0
        and metrics["calibration"]["state"] == "uncalibrated"
    )
    assert metrics["trends"][0]["rounds"][0]["open_qualification"] == 1

    # 正文改版 → 自动开新一轮；新正文满足工期要求，废标风险消失
    old_file = file_id_of(env.api, env.pid, body)
    _upload(env.api, store, env.pid, BODY_V2, name="施工组织设计.pdf", role="body", key="body-upload-0002", version_of=old_file)
    assert env.state() == "analyzing" and env.get("/pipeline")["stage"] == "running"
    env.drain()
    second = env.get("/report")
    assert second["round_no"] == 2 and second["snapshot_id"] != report["snapshot_id"]
    assert {c["category"]: c["total"] for c in second["categories"]}["disqualification"] == 0
    assert second["qualification"]["qualification"]["violated"] == 0
    trend = env.api.get("/metrics").json()["trends"][0]
    assert [r["round_no"] for r in trend["rounds"]] == [1, 2] and trend["rounds"][1]["open_qualification"] == 0

    # 重新评审是幂等的：没有失败环节、资料没变，不会重复入队
    with engine.connect() as conn:
        before = conn.exec_driver_sql("SELECT count(*) FROM jobs").scalar()
    again = env.api.post(f"/projects/{env.pid}/analysis")
    assert again.status_code == 202 and again.json()["stage"] == "done"
    with engine.connect() as conn:
        assert conn.exec_driver_sql("SELECT count(*) FROM jobs").scalar() == before


def test_failed_step_is_reported_and_restart_resumes(api_factory, make_member, store, engine, worker_container) -> None:  # type: ignore[no-untyped-def]
    llm = ScriptedLlm(_respond)
    env = _env(api_factory, make_member, store, engine, worker_container, llm)
    handlers = build_handlers(None)  # 先没有配置模型：规则提取明确失败
    env.parse_runner.handlers = handlers
    env.ai_runner.handlers = handlers
    _upload(env.api, store, env.pid, TENDER)
    _upload(env.api, store, env.pid, BODY, name="施工组织设计.pdf", role="body", key="body-upload-0001")
    env.drain()
    pipeline = env.get("/pipeline")
    assert pipeline["stage"] == "failed" and [s["state"] for s in pipeline["steps"]] == ["done", "failed", "waiting", "waiting"]
    assert env.state() == "analysis_failed"
    assert env.api.get(f"/projects/{env.pid}/report").json()["code"] == "report_not_ready"

    healthy = build_handlers(llm)  # 模型恢复后重新评审：失败的环节重试，其余不重做
    env.parse_runner.handlers = healthy
    env.ai_runner.handlers = healthy
    assert env.api.post(f"/projects/{env.pid}/analysis").status_code == 202
    env.drain()
    assert env.get("/pipeline")["stage"] == "done" and env.state() == "issues_open"
    with engine.connect() as conn:
        kinds = [r[0] for r in conn.exec_driver_sql("SELECT kind FROM jobs ORDER BY created_at")]
    assert kinds.count("parse") == 2 and kinds.count("rules") == 2  # 解析没有重做；规则提取失败一次、重试一次


def test_rule_edit_after_review_clones_package_and_asks_for_reanalysis(api_factory, make_member, store, engine, worker_container) -> None:  # type: ignore[no-untyped-def]
    env = _env(api_factory, make_member, store, engine, worker_container, ScriptedLlm(_respond))
    _upload(env.api, store, env.pid, TENDER)
    _upload(env.api, store, env.pid, BODY, name="施工组织设计.pdf", role="body", key="body-upload-0001")
    env.drain()
    rules = env.get("/rules")
    driver = next(r for r in rules["rules"] if r["kind"] == "disqualification")
    rejected = env.api.post(
        f"/rules/{driver['rule_id']}/reviews",
        {"object_version": str(rules["revision"]), "decision": "reject", "reason": "该条不属于本标段"},
    )
    assert rejected.status_code == 200, rejected.text
    now = env.get("/rules")
    assert now["rule_set_id"] != rules["rule_set_id"] and now["revision"] == rules["revision"] + 1  # 评审引用的规则包不变，更正落在新规则包
    assert env.get("/pipeline")["stale_reasons"] == ["rules_changed"]
    assert env.get("/report")["snapshot_id"] and any(a["key"] == "reanalyze" for a in env.get("/report")["next_actions"])
    env.drain()
    assert env.get("/report")["round_no"] == 1  # 规则更正不自动开新一轮，由人决定何时重新评审
    assert env.api.post(f"/projects/{env.pid}/analysis").status_code == 202
    env.drain()
    after = env.get("/report")
    assert after["round_no"] == 2 and after["qualification"]["qualification"]["total"] == 1
    assert all(r["statement"] != driver["statement"] for r in after["qualification"]["rows"])  # 被驳回的规则不再参与核对


def test_submission_and_outcome_feed_metrics_and_calibration(api_factory, make_member, store, engine, worker_container) -> None:  # type: ignore[no-untyped-def]
    env = _env(api_factory, make_member, store, engine, worker_container, ScriptedLlm(_respond))
    _upload(env.api, store, env.pid, TENDER)
    body = _upload(env.api, store, env.pid, BODY_V2, name="施工组织设计.pdf", role="body", key="body-upload-0001")[1].json()["version_id"]
    env.drain()
    report = env.get("/report")
    assert report["qualification"]["qualification"]["violated"] == 0  # 工期满足要求
    rules = env.get("/rules")
    for rule in rules["rules"]:
        if rule["kind"] in ("disqualification", "mandatory"):
            env.api.post(
                f"/rules/{rule['rule_id']}/reviews", {"object_version": str(rules["revision"]), "decision": "confirm", "reason": "核对原文"}
            )
    env.make_member("eng")
    reviewer = env.api_factory().login("eng")
    for finding in env.get("/findings?limit=100")["items"]:  # 余下的前后矛盾、范围残留等问题逐个处理
        reviewer.post(
            f"/findings/{finding['finding_id']}/reviews",
            {"object_version": str(finding["revision"]), "decision": "reject", "reason": "不成立"},
        )
    assert env.state() == "ready_to_submit" and any(a["key"] == "submit" for a in env.get("/report")["next_actions"])
    submit = env.api.post(
        f"/projects/{env.pid}/submissions",
        {
            "snapshot_id": report["snapshot_id"],
            "body_version_id": body,
            "submitted_at": "2026-10-01T08:00:00+08:00",
            "receipt_artifact_id": None,
        },
    )
    assert submit.status_code == 200, submit.text
    assert any(a["key"] == "record_outcome" for a in env.get("/report")["next_actions"])

    outcome = {
        "submission_id": submit.json()["submission_id"],
        "source_kind": "user_reported",
        "target": "technical_total",
        "value": "5",
        "maximum": "8",
        "aggregation": None,
        "stage_result": "technical_shortlisted",
        "source_page": None,
        "published_at": None,
        "supersedes": None,
        "source_artifact_id": None,
        "source_description": "开标现场公布",
    }
    assert env.api.post(f"/projects/{env.pid}/outcomes", outcome).status_code == 200
    metrics = env.api.get("/metrics").json()
    assert (
        metrics["outcomes"]["submissions"] == 1 and metrics["outcomes"]["shortlisted"] == 1 and metrics["outcomes"]["shortlist_rate"] == "1"
    )
    assert metrics["outcomes"]["award_rate"] is None  # 还没有最终结果：不给中标率
    calibration = metrics["calibration"]
    assert calibration["state"] == "insufficient" and calibration["pairs"] == 1 and calibration["mae"] is None  # 样本不足不给误差
    assert metrics["trends"][0]["actual_value"] == "5" and metrics["trends"][0]["maximum"] == "8"
    assert metrics["qualification"]["open_at_submission"] == 0
    # 同一提交后来又登记了结果（入围→中标）：校准只算一组，取最新一条
    later = {**outcome, "stage_result": "awarded", "value": "5.5"}
    assert env.api.post(f"/projects/{env.pid}/outcomes", later).status_code == 200
    again = env.api.get("/metrics").json()
    assert again["calibration"]["pairs"] == 1 and again["outcomes"]["awarded"] == 1 and again["outcomes"]["award_rate"] == "1"

    # 样本达到 20 组后才给出误差：直接补 19 组配对数据
    with engine.begin() as conn:
        conn.exec_driver_sql("SET LOCAL session_replication_role = replica")
        sub = conn.exec_driver_sql("SELECT id, project_id, snapshot_id FROM submissions").one()
        pred = conn.exec_driver_sql("SELECT estimated_decimal FROM predictions WHERE status='candidate'").scalar()
        for i in range(19):
            sid = f"sub_extra{i:04d}"
            conn.exec_driver_sql(
                "INSERT INTO submissions(id,project_id,body_version_id,snapshot_id,submitted_at,confirmation_review_id) "
                "SELECT %s,project_id,body_version_id,snapshot_id,submitted_at,confirmation_review_id FROM submissions WHERE id=%s",
                (sid, sub.id),
            )
            conn.exec_driver_sql(
                "INSERT INTO outcomes(id,project_id,submission_id,source_kind,verification_status,target,value_decimal,maximum_decimal,stage_result,known_at) "
                "VALUES (%s,%s,%s,'user_reported','pending','technical_total',5.5,8,'awarded',now())",
                (f"out_extra{i:04d}", sub.project_id, sid),
            )
    final = env.api.get("/metrics").json()["calibration"]
    assert final["state"] == "calibrated" and final["pairs"] == 20 and final["needed_pairs"] == 20
    expected_bias = (float(pred) / 8 - 5.5 / 8) * 100  # 估计 4 分对实际 5.5 分：偏低 18.75 个百分点
    assert float(final["bias"]) == pytest.approx(expected_bias, abs=0.01) and final["mae"] == final["bias"].lstrip("-")


def test_new_addendum_reextracts_rules_and_carries_decisions_into_next_round(
    api_factory, make_member, store, engine, worker_container
) -> None:  # type: ignore[no-untyped-def]
    env = _env(api_factory, make_member, store, engine, worker_container, ScriptedLlm(_respond))
    _upload(env.api, store, env.pid, TENDER)
    _upload(env.api, store, env.pid, BODY_V2, name="施工组织设计.pdf", role="body", key="body-upload-0001")
    env.drain()
    first_rules = env.get("/rules")
    quality = next(r for r in first_rules["rules"] if r["kind"] == "mandatory")
    ok = env.api.post(
        f"/rules/{quality['rule_id']}/reviews",
        {"object_version": str(first_rules["revision"]), "decision": "confirm", "reason": "核对原文"},
    )
    assert ok.status_code == 200, ok.text

    addendum = make_pdf([["补遗一", "评标办法和工期要求不变，其余按招标文件执行。"]])
    _upload(env.api, store, env.pid, addendum, name="补遗1.pdf", role="supplement", key="addendum-upload-0001")
    assert env.state() == "analyzing"
    env.drain()
    assert env.get("/report")["round_no"] == 2
    second_rules = env.get("/rules")
    assert second_rules["rule_set_id"] != first_rules["rule_set_id"] and len(second_rules["source_version_ids"]) == 2
    carried = next(r for r in second_rules["rules"] if r["statement"] == quality["statement"])
    assert (
        carried["confirmation"] == "confirmed" and carried["carried_from_rule_id"] == quality["rule_id"]
    )  # 招标文件改版后，同一条款沿用人工确认


def _set(env: Env, **settings: object) -> None:
    project = env.api.get(f"/projects/{env.pid}").json()
    r = env.api.patch(
        f"/projects/{env.pid}", {"settings": {**project["settings"], **settings}}, headers={"If-Match": str(project["revision"])}
    )
    assert r.status_code == 200, r.text


def test_manual_start_waits_for_the_user_each_round(api_factory, make_member, store, engine, worker_container) -> None:  # type: ignore[no-untyped-def]
    env = _env(api_factory, make_member, store, engine, worker_container, ScriptedLlm(_respond))
    assert env.api.get(f"/projects/{env.pid}").json()["settings"]["auto_start"] is True  # 默认自动开始
    _set(env, auto_start=False)
    _upload(env.api, store, env.pid, TENDER)
    body = _upload(env.api, store, env.pid, BODY, name="施工组织设计.pdf", role="body", key="body-upload-0001")[1].json()["version_id"]
    env.drain()
    pipeline = env.get("/pipeline")
    assert pipeline["stage"] == "waiting_start" and [s["state"] for s in pipeline["steps"]] == ["done", "waiting", "waiting", "waiting"]
    assert env.state() == "ready_to_analyze"  # 资料已解析（不调用模型），评审等人点“开始评审”
    with engine.connect() as conn:
        assert {r[0] for r in conn.exec_driver_sql("SELECT kind FROM jobs")} == {"parse"}

    assert env.api.post(f"/projects/{env.pid}/analysis").status_code == 202
    env.drain()
    assert env.get("/pipeline")["stage"] == "done" and env.state() == "issues_open"

    _upload(
        env.api,
        store,
        env.pid,
        BODY_V2,
        name="施工组织设计.pdf",
        role="body",
        key="body-upload-0002",
        version_of=file_id_of(env.api, env.pid, body),
    )
    env.drain()
    waiting = env.get("/pipeline")
    assert waiting["stage"] == "waiting_start" and waiting["stale_reasons"] == ["body_updated"]
    assert env.state() == "issues_open"  # 旧报告仍有效：状态回到它对应的状态，而不是停在“评审中”
    report = env.get("/report")
    assert report["round_no"] == 1 and report["next_actions"][0]["key"] == "start_analysis"  # 旧报告仍可看，并提示开始新一轮
    env.api.post(f"/projects/{env.pid}/analysis")
    env.drain()
    assert env.get("/report")["round_no"] == 2 and env.get("/pipeline")["stage"] == "done"


def test_switches_trim_checks_report_and_submission_gate(api_factory, make_member, store, engine, worker_container) -> None:  # type: ignore[no-untyped-def]
    env = _env(api_factory, make_member, store, engine, worker_container, ScriptedLlm(_respond))
    _set(env, consistency_check=False, scope_check=False, price=False, show_score=False, accept_rules_as_written=True)
    _upload(env.api, store, env.pid, TENDER)
    body = _upload(env.api, store, env.pid, BODY_V2, name="施工组织设计.pdf", role="body", key="body-upload-0001")[1].json()["version_id"]
    env.drain()
    report = env.get("/report")
    by = {c["category"]: c["total"] for c in report["categories"]}
    assert by["consistency"] == 0 and by["scope"] == 0  # 关闭的检测不产生问题
    assert report["price"] is None and [a["key"] for a in report["next_actions"]] == ["submit"]  # 不催报价、不催确认规则
    assert env.state() == "ready_to_submit"
    drawing = make_pdf([["图纸目录", "总平面图"]])
    _upload(env.api, store, env.pid, drawing, name="图纸.pdf", role="drawing", key="drawing-upload-0001")  # 非评审资料不影响提交
    submit = env.api.post(
        f"/projects/{env.pid}/submissions",
        {
            "snapshot_id": report["snapshot_id"],
            "body_version_id": body,
            "submitted_at": "2026-10-01T08:00:00+08:00",
            "receipt_artifact_id": None,
        },
    )
    assert submit.status_code == 200, submit.text  # 规则按原文接受：未确认的废标／强制类规则不再阻止登记
    docx = env.api.get(f"/projects/{env.pid}/report.docx")
    assert docx.status_code == 200
    from docx import Document

    text = "\n".join(p.text for p in Document(io.BytesIO(docx.content)).paragraphs)
    assert "可能的失分点" in text and "技术标估计" not in text and "四、报价" not in text and "按原文接受" in text
    with engine.connect() as conn:
        audit = conn.exec_driver_sql("SELECT details_json FROM audit_events WHERE action='project.update'").scalar()
    assert audit["settings"]["accept_rules_as_written"] is True  # 开关的改动留审计


def test_rules_extracted_from_outdated_tender_are_discarded(api_factory, make_member, store, engine, worker_container) -> None:  # type: ignore[no-untyped-def]
    """提取规则期间又上传了补遗：那次提取的结果不作为规则包，流水线针对最新资料重新提取，不会卡住。"""
    env = _env(api_factory, make_member, store, engine, worker_container, ScriptedLlm(_respond))
    _upload(env.api, store, env.pid, TENDER)
    _upload(env.api, store, env.pid, BODY_V2, name="施工组织设计.pdf", role="body", key="body-upload-0001")
    with engine.begin() as conn:
        tasks.relay_outbox(conn, env.bus.producer())
    assert env.parse_runner.poll_once() and env.parse_runner.poll_once()  # 两份资料解析完，规则提取入队但还没执行
    addendum = make_pdf([["补遗一", "评标办法和工期要求不变。"]])
    _upload(env.api, store, env.pid, addendum, name="补遗1.pdf", role="supplement", key="addendum-upload-0001")
    env.drain()
    assert env.get("/pipeline")["stage"] == "done" and env.state() in ("issues_open", "ready_to_submit")
    with engine.connect() as conn:
        packages = conn.exec_driver_sql("SELECT source_versions_json FROM rule_packages").all()
        rule_jobs = conn.exec_driver_sql("SELECT state FROM jobs WHERE kind='rules'").all()
    assert [len(p[0]) for p in packages] == [2]  # 只有基于招标文件＋补遗的规则包
    assert [j[0] for j in rule_jobs] == ["succeeded", "succeeded"]


def test_freeze_rejection_is_reported_and_does_not_block_others(api_factory, make_member, store, engine, worker_container) -> None:  # type: ignore[no-untyped-def]
    env = _env(api_factory, make_member, store, engine, worker_container, ScriptedLlm(_respond))
    project = env.api.get(f"/projects/{env.pid}").json()
    env.api.patch(f"/projects/{env.pid}", {"deadline": "2020-01-01T00:00:00+08:00"}, headers={"If-Match": str(project["revision"])})
    _upload(env.api, store, env.pid, TENDER)
    _upload(env.api, store, env.pid, BODY_V2, name="施工组织设计.pdf", role="body", key="body-upload-0001")
    env.drain()
    pipeline = env.get("/pipeline")
    review = pipeline["steps"][2]
    assert pipeline["stage"] == "failed" and review["state"] == "failed" and "截止" in review["message"]
    assert env.state() == "analysis_failed"
    from bidwb.services import pipeline as pipeline_service

    with engine.begin() as conn:  # 回收器兜底：出错的项目不拖垮整批
        conn.exec_driver_sql("UPDATE projects SET state='analyzing'")
        assert pipeline_service.advance_active(conn, store) == 1
        assert conn.exec_driver_sql("SELECT state FROM projects").scalar() == "analysis_failed"


def test_restart_reruns_incomplete_review_and_inputs_change_demotes_ready(
    api_factory, make_member, store, engine, worker_container
) -> None:  # type: ignore[no-untyped-def]
    broken = {"on": True}

    def respond(system: str, user: str) -> Any:
        if broken["on"] and "合规检查员" in system and "工程质量" in user.split("）：", maxsplit=1)[-1][:10]:
            return "这不是 JSON"
        return _respond(system, user)

    env = _env(api_factory, make_member, store, engine, worker_container, ScriptedLlm(respond))
    _set(env, consistency_check=False, scope_check=False, accept_rules_as_written=True)
    _upload(env.api, store, env.pid, TENDER)
    _upload(env.api, store, env.pid, BODY_V2, name="施工组织设计.pdf", role="body", key="body-upload-0001")
    env.drain()
    assert env.get("/report")["basis"]["coverage"]["status"] == "partial" and env.state() == "issues_open"
    broken["on"] = False
    assert env.api.post(f"/projects/{env.pid}/analysis").status_code == 202  # 覆盖不完整的核对可以重新跑
    env.drain()
    assert env.get("/report")["basis"]["coverage"]["status"] == "complete" and env.state() == "ready_to_submit"

    rules = env.get("/rules")
    scoring = next(r for r in rules["rules"] if r["kind"] == "scoring")
    env.api.post(
        f"/rules/{scoring['rule_id']}/reviews", {"object_version": str(rules["revision"]), "decision": "reject", "reason": "不适用"}
    )
    assert env.state() == "issues_open"  # 规则已更正、报告落后于输入：不能停在“可登记提交”
    assert all(a["key"] != "submit" for a in env.get("/report")["next_actions"])


def test_switching_auto_start_on_starts_the_review(api_factory, make_member, store, engine, worker_container) -> None:  # type: ignore[no-untyped-def]
    env = _env(api_factory, make_member, store, engine, worker_container, ScriptedLlm(_respond))
    _set(env, auto_start=False)
    _upload(env.api, store, env.pid, TENDER)
    _upload(env.api, store, env.pid, BODY_V2, name="施工组织设计.pdf", role="body", key="body-upload-0001")
    env.drain()
    assert env.state() == "ready_to_analyze"
    _set(env, auto_start=True)
    env.drain()
    assert env.get("/pipeline")["stage"] == "done"


def test_oversized_price_numbers_are_rejected_not_crashing(api_factory, make_member, store, engine, worker_container) -> None:  # type: ignore[no-untyped-def]
    env = _env(api_factory, make_member, store, engine, worker_container, ScriptedLlm(_respond))
    params = {
        "full_score": "30",
        "above_deduct_per_pct": "1",
        "below_deduct_per_pct": "0.5",
        "benchmark": {"method": "average", "coefficient": "1", "trim_pct": "0", "include_own": False},
    }
    big = {"params": params, "competitor_prices": ["1"], "candidate_prices": ["9" * 24], "note": ""}
    ok = env.api.post(f"/projects/{env.pid}/price:simulate", big)
    assert ok.status_code == 200, ok.text  # 24 位的极端报价也能算，不溢出
    too_big = env.api.post(f"/projects/{env.pid}/price:simulate", {**big, "candidate_prices": ["9" * 25]})
    assert too_big.status_code == 422


def test_bid_terms_feed_coverage_price_template_and_overview(api_factory, make_member, store, engine, worker_container) -> None:  # type: ignore[no-untyped-def]
    env = _env(api_factory, make_member, store, engine, worker_container, ScriptedLlm(_respond))
    overview = env.api.get("/projects").json()[0]["summary"]
    assert overview["reviewed"] is False and overview["next_action"]["key"] == "upload_materials"
    assert "招标文件" in overview["next_action"]["title"]
    _upload(env.api, store, env.pid, TENDER)
    _upload(env.api, store, env.pid, BODY, name="施工组织设计.pdf", role="body", key="body-upload-0001")
    env.drain()

    terms = env.get("/rules")["terms"]
    assert terms["technical_full_score"]["value"] == "60" and terms["technical_full_score"]["evidence"][0]["file_role"] == "tender"
    price_rule = terms["price_rule"]
    assert (price_rule["full_score"], price_rule["above_deduct_per_pct"], price_rule["below_deduct_per_pct"]) == ("30", "1", "0.5")
    assert price_rule["max_limit"] is None  # 原文里没有的数值被丢弃
    prediction = env.get("/report")["prediction"]
    assert (prediction["maximum"], prediction["technical_full_score"]) == ("8", "60")  # 只覆盖了 8／60 分

    row = env.api.get("/projects").json()[0]["summary"]
    report = env.get("/report")
    assert row["reviewed"] is True and row["qualification_open"] == 1
    assert row["next_action"] == report["next_actions"][0]  # 总览与报告同源
    assert row["prediction"]["technical_full_score"] == "60" and row["prediction"]["estimated_value"] == "4"

    # 批量确认未审核的规则：已有决定的不受影响，不复制规则包
    rules = env.get("/rules")
    first = rules["rules"][0]["rule_id"]
    env.api.post(f"/rules/{first}/reviews", {"object_version": str(rules["revision"]), "decision": "reject", "reason": "不适用"})
    rules = env.get("/rules")
    ids = [r["rule_id"] for r in rules["rules"]]
    stale = env.api.post(f"/projects/{env.pid}/rules:confirm", {"object_version": "999", "rule_ids": ids})
    assert stale.status_code == 412
    done = env.api.post(f"/projects/{env.pid}/rules:confirm", {"object_version": str(rules["revision"]), "rule_ids": ids})
    assert done.status_code == 200 and done.json() == {"confirmed": len(ids) - 1, "skipped": 1}
    after = env.get("/rules")
    assert [r["confirmation"] for r in after["rules"]].count("confirmed") == len(ids) - 1
    unknown = env.api.post(
        f"/projects/{env.pid}/rules:confirm", {"object_version": str(after["revision"]), "rule_ids": ["rul_missing0001"]}
    )
    assert unknown.status_code == 404
